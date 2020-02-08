import time
import requests
import re
import openpyxl
import threading


headers = {'user-agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/79.0.3945.130 Safari/537.36',
           'ctaapi-version': '1.1'}
wb = openpyxl.load_workbook(filename='books.xlsx')
sheet = wb['Лист1']


def write_to_excel(row, *tuple_of_data):
    """
        Функция записи в xlsx файл:
    :param row: передаваемый ряд (какая по счёту запись или же компания)
    :param tuple_of_data:   данные о самой компании, если они есть
    :return: ничего не возвращает, только сохраняет изменяемый файл
    """
    # try:
    #     row = len(list(sheet))
    # except TypeError:
    #     row = 1
    if row == 1:    # если первая запись то пишет вначале поля
        name_of_coloms = ('Название', 'Описание', 'Адрес', 'Номер', 'Ссылка', 'Booth')  # сами поля
        for i in enumerate(name_of_coloms):
            sheet.cell(row, i[0] + 1, i[1])
    row += 1  # увеличиваем ряд на 1, из-за полей
    for i in enumerate(tuple_of_data):
        sheet.cell(row, i[0] + 1, i[1])
    wb.save('books.xlsx')


def start_download(page):
    html = requests.get(
        'https://www.ces.tech/api/Exhibitors?searchTerm=&sortBy=alpha&filter=&pageNo=' + str(page) + '&pageSize=30',
        headers=headers)
    if html.status_code != 200:
        quit()
    # print(html.text)
    html = html.text[html.text.find('"exhibitors":') + 13:]
    while html.find('{"companyName"') > -1:
        html = html[html.find('{"companyName"'):]
        print(html)
        current_company = html[:html.find(',{"companyName":"')]  # теперь current company - эта вся рассматриваемая компания\
        name_company = re.search(r'{\"companyName\":.*"alpha"', current_company).group()[16:-9]
        html = html[15:]
        print(name_company)
        description_company = re.search(r'\"description\":[^:]*', current_company).group()[15:-15]
        companyLink = re.search(r'\"companyLink\":.*\",\"isCategorySponsor\":', current_company).group()[15:-22]
        about_company = requests.get(companyLink)  # html страница компании (для того чтоб парсить больше инфы)
        if about_company.status_code != 200:
            html = html[html.find('{"companyName"'):]
            continue
        about_company = about_company.text
        address = about_company[
                  about_company.find('<div class="dtc  pa0  pl2">') + 27: about_company.find('<p class="mb0">')]
        address = ''.join(tuple(x if not x.isspace() else '' for x in re.sub(r'(/?<br.*/>)', '', address)))
        address += ' ' + re.search(r'<span class="break-word  lh-list">.*</span>', about_company).group()[34:-7]
        phone = booth = url = ''
        if about_company.find('<span class="break-word  lh-list  muted  pr2">(p):</span>') > -1:
            phone = about_company[about_company.find('<span class="break-word  lh-list  muted  pr2">(p):</span>') + 56:]
            phone = re.search(r'>.*<', phone).group()[1:-1]

        if about_company.find('<p class="lh-copy  ma0">') > -1:
            url = about_company[about_company.find('<p class="lh-copy  ma0">') + 24:]
            url = re.search(r'href=\".*\"', url[:url.find('</a')]).group()[6:-1]

        if about_company.find('class="mys-floorPlanLink" style="">') > -1:
            booth = about_company[about_company.find('class="mys-floorPlanLink" style="">') + 35:]
            booth = booth[:booth.find('</a>')]
        # print(i, name_company, description_company, address, phone, url, booth)
        write_to_excel(page, name_company, description_company, address, phone, url, booth)


if __name__ == '__main__':
    page = 1  # страница по счёту которую мы парсим
    # print(requests.get('https://www.ces.tech/api/Exhibitors?searchTerm=&sortBy=alpha&filter=&pageNo=5&pageSize=30',
    #       headers=headers).text)
    # quit()
    while True:
        # threading.Thread(target=start_download, args=(page,)).start()
        # time.sleep(1)
        start_download(page)
        quit()
        # quit()
        # time.sleep(3)
        # if page % 50 == 0:
        #     time.sleep(5)
        #     quit()
        page += 1

